import datetime
import os
import docx
from num2words import num2words
from sqlalchemy.orm import Session
from fastapi import HTTPException
import locale

# Configurar el locale para que los nombres de los meses salgan en español
try:
    locale.setlocale(locale.LC_TIME, 'es_ES.UTF-8')
except locale.Error:
    try:
        locale.setlocale(locale.LC_TIME, 'Spanish_Spain')
    except locale.Error:
        print("Advertencia: No se pudo configurar el locale en español.")

from app.repository.empresa import EmpresaRepository
from app.repository.representante import RepresentanteRepository
from app.models.documento import GenerationRequest
from app.utils.ayudante_docx import reemplazar_placeholders_docx

class ServicioDocumento:
    def __init__(self):
        self.repo_empresa = EmpresaRepository()
        self.repo_representante = RepresentanteRepository()

    def formato_fecha_largo(self, objeto_fecha):
        if not isinstance(objeto_fecha, (datetime.date, datetime.datetime)):
            return ""
        day_num = objeto_fecha.day
        day_letras = num2words(day_num, lang='es')
        month_name = objeto_fecha.strftime('%B')
        return f"el {day_letras} ({day_num}) de {month_name} de {objeto_fecha.year}"

    def generar_documento(self, db: Session, solicitud: GenerationRequest):
        # --- SINTAXIS CORREGIDA A SNAKE_CASE ---
        resultado_empresa = self.repo_empresa.get_by_id(db, solicitud.empresa_id)
        if not resultado_empresa:
            raise HTTPException(status_code=404, detail="Empresa no encontrada")

        resultado_representante = self.repo_representante.get_by_id(db, solicitud.representante_id)
        if not resultado_representante:
            raise HTTPException(status_code=404, detail="Representante no encontrado")

        # --- PREPARACIÓN DE DATOS ---
        rep_edad_num = (datetime.date.today() - resultado_representante['fecha_nacimiento']).days // 365
        rep_edad_letras = num2words(rep_edad_num, lang='es').upper()
        rep_cui_letras = " ".join(num2words(c, lang='es') for c in resultado_representante['cui']).upper()
        
        genero_texto = "El Notario" # Valor por defecto

        num_registro_letras = num2words(int(resultado_empresa['numero_registro']), lang='es').upper() if resultado_empresa.get('numero_registro') and resultado_empresa['numero_registro'].isdigit() else resultado_empresa.get('numero_registro', '')
        num_libro_letras = num2words(int(resultado_empresa['numero_libro']), lang='es').upper() if resultado_empresa.get('numero_libro') and resultado_empresa['numero_libro'].isdigit() else resultado_empresa.get('numero_libro', '')
        num_folio_letras = num2words(int(resultado_empresa['numero_folio']), lang='es').upper() if resultado_empresa.get('numero_folio') and resultado_empresa['numero_folio'].isdigit() else resultado_empresa.get('numero_folio', '')
        
        colaborador = solicitud.colaborador_data.datos_persona
        colab_cui_letras = " ".join(num2words(c, lang='es') for c in colaborador.cui).upper() if colaborador.cui and colaborador.cui.isdigit() else ""
        colab_edad_letras = num2words(int(colaborador.edad), lang='es').upper() if colaborador.edad and colaborador.edad.isdigit() else ""
        colab_estado_civil = colaborador.estado_civil or "soltero"
        colab_nacionalidad = colaborador.nacionalidad or "guatemalteco"
        colab_profesion = colaborador.profesion or "N/A"
        colab_posicion = colaborador.posicion or solicitud.colaborador_data.datos_contrato.tipo_contrato

        try:
            fecha_inicio = datetime.datetime.strptime(solicitud.colaborador_data.datos_contrato.fecha_inicio, "%Y-%m-%d")
            dia_inicio, mes_inicio, anio_inicio = fecha_inicio.day, fecha_inicio.strftime('%B'), fecha_inicio.year
            dia_inicio_letras, anio_inicio_letras = num2words(dia_inicio, lang='es'), num2words(anio_inicio, lang='es')
        except (ValueError, TypeError):
            dia_inicio, mes_inicio, anio_inicio, dia_inicio_letras, anio_inicio_letras = "N/A", "N/A", "N/A", "N/A", "N/A"

        try:
            fecha_fin_str = solicitud.colaborador_data.datos_contrato.fecha_fin
            if fecha_fin_str and "indefinido" not in fecha_fin_str.lower():
                fecha_fin = datetime.datetime.strptime(fecha_fin_str, "%Y-%m-%d")
                dia_fin, mes_fin, anio_fin = fecha_fin.day, fecha_fin.strftime('%B'), fecha_fin.year
                dia_fin_letras, anio_fin_letras = num2words(dia_fin, lang='es'), num2words(anio_fin, lang='es')
            else:
                dia_fin, mes_fin, anio_fin, dia_fin_letras, anio_fin_letras = "N/A", "N/A", "N/A", "N/A", "N/A"
        except (ValueError, TypeError):
            dia_fin, mes_fin, anio_fin, dia_fin_letras, anio_fin_letras = "N/A", "N/A", "N/A", "N/A", "N/A"

        reemplazos = {
            '{{nombre_completo}}': colaborador.nombre_completo,
            '{{cui}}': colaborador.cui,
            '{{cui_letras}}': colab_cui_letras,
            '{{edad_empleado}}': colaborador.edad,
            '{{edad_empleado_letras}}': colab_edad_letras,
            '{{direccion}}': colaborador.direccion,
            '{{estado_civil}}': colab_estado_civil,
            '{{nacionalidad}}': colab_nacionalidad,
            '{{profesion}}': colab_profesion,
            '{{posicion}}': colab_posicion,
            '{{puesto}}': colab_posicion,
            '{{colaborador_lugar_notificaciones}}': colaborador.direccion,
            '{{fecha_contrato}}': solicitud.fecha_contrato,
            '{{día_letras}}': dia_inicio_letras,
            '{{día_numeros}}': str(dia_inicio),
            '{{mes_letras}}': mes_inicio,
            '{{año_letras}}': anio_inicio_letras,
            '{{año_numeros}}': str(anio_inicio),
            '{{vence_dia_letreas}}': dia_fin_letras,
            '{{vence_dia_numeros}}': str(dia_fin),
            '{{vence_mes_letras}}': mes_fin,
            '{{vence_año_letras}}': anio_fin_letras,
            '{{vence_año_numeros}}': str(anio_fin),
            '{{monto}}': solicitud.colaborador_data.datos_contrato.monto,
            '{{monto_letras}}': solicitud.colaborador_data.datos_contrato.monto_en_letras,
            '{{rep_legal_nombre}}': resultado_representante['nombre_completo'],
            '{{rep_legal_edad}}': str(rep_edad_num),
            '{{rep_legal_edad_letras}}': rep_edad_letras,
            '{{rep_legal_estado_civil}}': resultado_representante['estado_civil'],
            '{{rep_legal_profesion}}': resultado_representante['profesion'],
            '{{rep_legal_nacionalidad}}': resultado_representante['nacionalidad'],
            '{{rep_legal_cui}}': resultado_representante['cui'],
            '{{rep_legal_cui_letras}}': rep_cui_letras,
            '{{rep_legal_extendido_en}}': resultado_representante['extendido_en'],
            '{{genero}}': genero_texto,
            '{{empresa_contratante}}': resultado_empresa['razon_social'],
            '{{empresa_entidad}}': resultado_empresa['razon_social'],
            '{{empresa_autorizada_en}}': resultado_empresa['autorizada_en'],
            '{{empresa_fecha_autorizacion}}': self.formato_fecha_largo(resultado_empresa['fecha_autorizacion']),
            '{{empresa_autorizada_por}}': resultado_empresa['autorizada_por'],
            '{{empresa_inscrita_en}}': resultado_empresa['inscrita_en'],
            '{{empresa_numero_registro}}': resultado_empresa['numero_registro'],
            '{{empresa_numero_registro_letras}}': num_registro_letras,
            '{{empresa_numero_folio}}': resultado_empresa['numero_folio'],
            '{{empresa_numero_folio_letras}}': num_folio_letras,
            '{{empresa_numero_libro}}': resultado_empresa['numero_libro'],
            '{{empresa_numero_libro_letras}}': num_libro_letras,
            '{{empresa_tipo_libro}}': resultado_empresa['tipo_libro'],
            '{{empresa_lugar_notificaciones}}': resultado_empresa['lugar_notificaciones'],
            '{{empresa_segundo_lugar_notificaciones}}': resultado_empresa['segundo_lugar_notificaciones'],
        }

        ruta_plantilla = os.path.join("templates", solicitud.template_name)
        if not os.path.exists(ruta_plantilla):
             raise HTTPException(status_code=404, detail=f"Plantilla no encontrada: {ruta_plantilla}")
        
        doc = docx.Document(ruta_plantilla)
        reemplazar_placeholders_docx(doc, reemplazos)
        
        nombre_archivo_salida = f"contrato_generado_{colaborador.nombre_completo.replace(' ', '_') if colaborador.nombre_completo else 'sin_nombre'}.docx"
        doc.save(nombre_archivo_salida)
        
        return nombre_archivo_salida

